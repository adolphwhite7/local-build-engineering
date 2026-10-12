"""Read-only original XLSX fields and keyed joins; cached formulas are separate."""
import datetime
import hashlib
import re
from thank_you_discovery import entity_candidates

KEYS={'Gifts':'Gift ID','Giver Links':'Link ID','Guests':'Guest ID','Households':'Household ID','Thank You':'Thank-you ID'}

def value(cell):return cell['value']
def iso(v):return v.date().isoformat() if isinstance(v,datetime.datetime) else v.isoformat() if isinstance(v,datetime.date) else v

def read(path):
    """Read original and data-only cells without saving, evaluating or network."""
    import openpyxl  # Optional workbook adapter dependency; paging uses only the standard library.
    before=hashlib.sha256(path.read_bytes()).hexdigest()
    w=openpyxl.load_workbook(path,read_only=True,data_only=False)
    cached=openpyxl.load_workbook(path,read_only=True,data_only=True)
    tables={};summary={};owners={};metadata={}
    try:
        for name,key in KEYS.items():
            sheet=w[name];rows=list(sheet.iter_rows());cache=list(cached[name].iter_rows())
            header=next(i for i,r in enumerate(rows) if any(c.value==key for c in r))
            headers=[c.value for c in rows[header]]
            if len([h for h in headers if h is not None])!=len(set(h for h in headers if h is not None)):raise ValueError('Duplicate native headers')
            records=[]
            for row,cr in zip(rows[header+1:],cache[header+1:]):
                item={h:{'value':iso(c.value),'cell':name+'!'+c.coordinate,
                         'formula':c.value if c.data_type=='f' else None,
                         'cached':iso(cc.value) if c.data_type=='f' else None}
                      for h,c,cc in zip(headers,row,cr) if h is not None}
                if item.get(key,{}).get('value') is not None:records.append(item)
            tables[name]=records
            metadata[name]=[str(c.value) for row in rows[:header] for c in row if c.value is not None]
        # Labeled summary inputs and declared owner keys, not ordinal record selection.
        for row in w['Summary'].iter_rows():
            for left,right in zip(row,row[1:]):
                if left.value in ('Wedding date','Early target days','Event or later days','Hard backstop','Early rule approved','Gift ledger revision'):
                    summary[left.value]={'value':iso(right.value),'cell':'Summary!'+right.coordinate}
                if isinstance(left.value,str) and re.fullmatch(r'P\d+',left.value):
                    if left.value in owners:raise ValueError('Duplicate owner ID')
                    owners[left.value]={'value':right.value,'cell':'Summary!'+right.coordinate}
    finally:w.close();cached.close()
    if hashlib.sha256(path.read_bytes()).hexdigest()!=before:raise ValueError('Native source changed')
    return {'path':str(path),'sha256':before,'tables':tables,'summary_inputs':summary,'owners':owners,'metadata':metadata,
            'formula_handling':'Formula text and saved cached values read separately; no workbook recalculation or edit performed.'}

def joined(request,native):
    t=native['tables'];candidates=entity_candidates(request,[{'id':value(r['Gift ID']),'label':value(r['Gift'])} for r in t['Gifts']])
    result={'request':request,'candidates':[{'id':k,'label':v} for k,v in candidates],'unknown':[], 'source_cells':{},'model_reasoning_credit':False}
    if len(candidates)!=1:result['status']='ENTITY_MISSING_OR_AMBIGUOUS';return result
    fid,label=candidates[0]
    result['fields']={'gift':label,'gift_id':fid}
    def one(sheet,key,v):
        rows=[r for r in t[sheet] if value(r[key])==v]
        if len(rows)!=1:raise ValueError(sheet+' keyed match missing or ambiguous')
        return rows[0]
    try:
        gift=one('Gifts','Gift ID',fid);note=one('Thank You','Gift ID',fid)
        if value(gift['Thank-you ID'])!=value(note['Thank-you ID']):raise ValueError('Reciprocal gift/workflow ID mismatch')
        links=[r for r in t['Giver Links'] if value(r['Gift ID'])==fid]
        if not links:raise ValueError('Giver link missing')
        if any(value(r['Giver group ID'])!=value(note['Giver group ID']) for r in links):raise ValueError('Gift group mismatch')
        result['recipients']=[]
        for link in links:
            h=one('Households','Household ID',value(link['Household ID']))
            pid=value(link['Person ID']);g=one('Guests','Person ID',pid) if pid else None
            if g and (value(g['Household ID'])!=value(h['Household ID']) or value(g['Guest ID']) not in [s.strip() for s in str(value(h['Guest IDs'])).split(',')]):raise ValueError('Person household membership mismatch')
            result['recipients'].append({'card_credit':value(link['Name credited on card']),
                'person_id':pid,'household_id':value(h['Household ID']),
                'relationship':value(g['Relationship']) if g else None,'personal_note':value(g['Personal note']) if g else None,
                'attendance':value(g['Attendance (simulated)']) if g else None,
                'address_publication':value(h['Publish address?']),'deliverable':value(h['Deliverable?']),
                'recipient_note':value(link['Recipient note wording']),'recipient_approval':value(link['Partner approval']),
                'recipient_sent':value(link['Recipient sent (simulated)'])})
            for sheet,row in [('Giver Links',link),('Households',h),('Guests',g)]:
                if row:
                    for k,cell in row.items():result['source_cells'][cell['cell']]=dict(cell,header=k)
        for sheet,row in [('Gifts',gift),('Thank You',note)]:
            for k,cell in row.items():result['source_cells'][cell['cell']]=dict(cell,header=k)
        owner=value(note['Owner ID']);owner_record=native['owners'].get(owner)
        result['fields']={'gift':label,'gift_id':fid,'correspondence_id':value(note['Thank-you ID']),
            'group_id':value(note['Giver group ID']),'credited_giver':'; '.join(str(r['card_credit']) for r in result['recipients'] if r['card_credit']),
            'purchaser_id':value(gift['Purchaser ID']),'purchaser_information':value(gift['Purchaser information']),
            'received':value(gift['Received']),'owner_id':owner,'owner_name':owner_record['value'] if owner_record else None,
            'status_record':value(note['Status']),'policy_days':value(note['Recorded target days']),
            'recorded_target':value(note['Recorded due']),'draft_date':value(note['Draft date']),
            'approval_date':value(note['Partner approval']),'sent_date':value(note['Sent date (simulated)']),
            'note_wording':value(note['Note wording']),'recorded_next_action':value(note['Recorded next action'])}
        f=result['fields'];policy=native['summary_inputs']
        received=datetime.date.fromisoformat(f['received']);wedding=datetime.date.fromisoformat(policy['Wedding date']['value'])
        days=policy['Early target days' if received<wedding else 'Event or later days']['value']
        if type(days)!=int or not 0<days<=365:raise ValueError('Native policy interval invalid')
        cap=datetime.date.fromisoformat(policy['Hard backstop']['value'])
        computed=min(received+datetime.timedelta(days=days),cap).isoformat()
        result['date_calculation']={'computed_due':computed,'days':days,'backstop':cap.isoformat(),
            'matches_recorded':computed==f['recorded_target'] and days==f['policy_days'],
            'cached_due':value(note['Calculated due']) if not note['Calculated due']['formula'] else note['Calculated due']['cached'],
            'formula':note['Calculated due']['formula'],'cache_is_not_recalculation':True,
            'policy_approved':policy['Early rule approved']['value'],
            'scope':'Current workbook retrospective rule; approval date not moved back to earlier receipt.'}
        if not result['date_calculation']['matches_recorded']:result['unknown'].append('Calculated/recorded date or interval disagreement; target unresolved')
        for key in ['purchaser_id','approval_date','sent_date','note_wording','owner_name']:
            if f[key] is None:result['unknown'].append(key+' not recorded; blank is not an explicit negative')
        for recipient in result['recipients']:
            if recipient['personal_note'] is None:result['unknown'].append('Personal note/individual memories not recorded; do not invent wording')
        result['policy_sources']=policy;result['owner_source']=owner_record
        result['status']='USEFUL_PARTIAL_NATIVE_WORK_ITEM'
    except (ValueError,KeyError,TypeError) as error:
        result['status']='PARTIAL_NATIVE_RELATIONSHIP_OR_DATE_UNRESOLVED';result['unknown'].append(str(error))
    return result
